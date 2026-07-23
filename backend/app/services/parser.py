"""
Handles extracting raw text from an uploaded resume file (.pdf or .docx).
"""

import io

import docx
import pdfplumber
from fastapi import HTTPException


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_chunks = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_chunks.append(page_text)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read PDF file. It may be corrupted or password protected. ({exc})",
        )

    text = "\n".join(text_chunks).strip()
    if not text:
        raise HTTPException(
            status_code=400,
            detail="No readable text found in the PDF. If this is a scanned "
            "image-based resume, please upload a text-based PDF or DOCX instead.",
        )
    return text


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read DOCX file. It may be corrupted. ({exc})",
        )

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Also pull text out of any tables in the document
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    text = "\n".join(paragraphs).strip()
    if not text:
        raise HTTPException(
            status_code=400,
            detail="No readable text found in the DOCX file.",
        )
    return text


def extract_resume_text(filename: str, content_type: str, file_bytes: bytes) -> str:
    """
    Dispatch to the correct extractor based on file extension / content type.
    """
    name = (filename or "").lower()

    is_pdf = name.endswith(".pdf") or content_type == "application/pdf"
    is_docx = name.endswith(".docx") or content_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    if is_pdf:
        return extract_text_from_pdf(file_bytes)
    if is_docx:
        return extract_text_from_docx(file_bytes)

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Please upload a PDF or DOCX resume.",
    )
